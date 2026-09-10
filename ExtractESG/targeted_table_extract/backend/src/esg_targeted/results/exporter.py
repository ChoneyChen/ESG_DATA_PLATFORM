from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from esg_standard_packages.contracts import CompiledStandardPackage

from esg_targeted.io import write_json, write_jsonl
from esg_targeted.results.spec import EXPORT_FILES, RECORD_FILES, RESULT_SCHEMA_VERSION
from esg_targeted.results.organization import FactOrganizer, element_order


SHEET_NAMES = {
    "reporting_tasks": "任务",
    "quantitative_observations": "量化事实",
    "qualitative_assertions": "定性断言",
    "attribute_values": "属性",
    "dimension_values": "维度",
    "evidence_references": "证据",
}


class ResultExporter:
    def export(
        self,
        *,
        output_dir: Path,
        records: dict[str, list[dict[str, Any]]],
        summary: dict[str, Any],
        package: CompiledStandardPackage,
    ) -> dict[str, str]:
        records, organization = FactOrganizer().organize(records, package)
        result_dir = output_dir / "results"
        export_dir = output_dir / "exports"
        result_dir.mkdir(parents=True, exist_ok=True)
        export_dir.mkdir(parents=True, exist_ok=True)
        for key, filename in RECORD_FILES.items():
            write_jsonl(result_dir / filename, records.get(key, []))
        write_json(result_dir / "summary.json", summary)
        write_json(result_dir / "fact-organization.json", organization)

        workbook_path = output_dir / EXPORT_FILES["workbook"]
        self._write_workbook(workbook_path, records, summary, package)
        task_csv = output_dir / EXPORT_FILES["task_csv"]
        self._write_csv(task_csv, records.get("reporting_tasks", []))
        self._write_csv(output_dir / EXPORT_FILES["fact_csv"], self._machine_fill_rows(records, summary, package))
        package_path = output_dir / EXPORT_FILES["result_package"]
        write_json(
            package_path,
            {
                "schema_version": RESULT_SCHEMA_VERSION,
                "job_id": summary["job_id"],
                "core_schema": summary["core_schema"],
                "summary": summary,
                "records": records,
                "organization": organization,
            },
        )
        return dict(EXPORT_FILES)

    def _write_workbook(
        self,
        path: Path,
        records: dict[str, list[dict[str, Any]]],
        summary: dict[str, Any],
        package: CompiledStandardPackage,
    ) -> None:
        workbook = Workbook()
        workbook.remove(workbook.active)
        machine_sheet = workbook.create_sheet("机器填写行")
        machine_rows = self._machine_fill_rows(records, summary, package)
        machine_headers = self._machine_fill_headers(package)
        machine_sheet.append(machine_headers)
        for row in machine_rows:
            machine_sheet.append(
                [self._cell_value(row.get(header)) for header in machine_headers]
            )
        self._style_sheet(machine_sheet)

        summary_sheet = workbook.create_sheet("运行摘要")
        summary_sheet.append(["field", "value"])
        for key, value in summary.items():
            summary_sheet.append([key, self._cell_value(value)])
        self._style_sheet(summary_sheet)

        for key, sheet_name in SHEET_NAMES.items():
            sheet = workbook.create_sheet(sheet_name)
            rows = records.get(key, [])
            headers = self._headers(rows)
            if headers:
                sheet.append(headers)
                for row in rows:
                    sheet.append([self._cell_value(row.get(header)) for header in headers])
            self._style_sheet(sheet)
        workbook.save(path)

    @staticmethod
    def _machine_fill_headers(package: CompiledStandardPackage) -> list[str]:
        headers = [
            "task_id",
            "metric_id",
            "source_datapoint_id",
            "metric_label",
            "fact_id",
            "fact_type",
            "found_status",
            "review_status",
        ]
        for element in sorted(package.elements, key=lambda e: element_order(e.model_dump(mode="json"))):
            if element.element_code not in headers:
                headers.append(element.element_code)
        return [
            *headers,
            "evidence_pdf_pages",
            "evidence_excerpts",
            "evidence_ids",
            "logical_measurement_id",
            "presentation_relation",
        ]

    @classmethod
    def _machine_fill_rows(
        cls,
        records: dict[str, list[dict[str, Any]]],
        summary: dict[str, Any],
        package: CompiledStandardPackage,
    ) -> list[dict[str, Any]]:
        metrics = {item.metric_id: item for item in package.metrics}
        elements = {item.element_id: item for item in package.elements}
        attributes = cls._rows_by_parent(records.get("attribute_values", []))
        dimensions = cls._rows_by_parent(records.get("dimension_values", []))
        evidence = cls._rows_by_parent(
            records.get("evidence_references", []), key="source_record_id"
        )
        facts_by_task: dict[str, list[tuple[str, dict[str, Any]]]] = {}
        for fact_type, collection, primary_key in (
            ("quantitative_observation", "quantitative_observations", "observation_id"),
            ("qualitative_assertion", "qualitative_assertions", "assertion_id"),
        ):
            for fact in records.get(collection, []):
                facts_by_task.setdefault(fact["task_id"], []).append((fact_type, fact))

        result = []
        _, organization = FactOrganizer().organize(records, package)
        for task in records.get("reporting_tasks", []):
            metric = metrics.get(task["metric_id"])
            task_facts = facts_by_task.get(task["task_id"]) or [("", {})]
            for fact_type, fact in task_facts:
                fact_id = fact.get("observation_id") or fact.get("assertion_id")
                row = {
                    "task_id": task["task_id"],
                    "metric_id": task["metric_id"],
                    "source_datapoint_id": (
                        metric.source_datapoint_id if metric else task["metric_id"]
                    ),
                    "metric_label": (
                        metric.labels.get("zh") or metric.labels.get("en")
                        if metric
                        else task["metric_id"]
                    ),
                    "fact_id": fact_id,
                    "fact_type": fact_type,
                    "found_status": task.get("found_status"),
                    "review_status": fact.get("review_status") or task.get("review_status"),
                }
                linked_values = {
                    item.get("element_id"): item
                    for item in [
                        *attributes.get(fact_id, []),
                        *dimensions.get(fact_id, []),
                    ]
                }
                for element_id in package.elements_by_metric.get(task["metric_id"], []):
                    element = elements[element_id]
                    fixed = element.value_contract.fixed_value
                    if fixed is not None:
                        value = fixed
                    elif element.binding.storage.value == "core_field":
                        value = fact.get(element.binding.target)
                    else:
                        value = cls._linked_value(linked_values.get(element_id))
                    row[element.element_code] = value
                citations = evidence.get(fact_id, [])
                row["evidence_pdf_pages"] = sorted(
                    {
                        int(item["pdf_page_index"]) + 1
                        for item in citations
                        if item.get("pdf_page_index") is not None
                    }
                )
                row["evidence_excerpts"] = [
                    item.get("excerpt") for item in citations if item.get("excerpt")
                ]
                row["evidence_ids"] = [
                    item.get("evidence_id") for item in citations if item.get("evidence_id")
                ]
                organized = organization["facts"].get(fact_id, {})
                row["logical_measurement_id"] = organized.get("logical_measurement_id")
                row["presentation_relation"] = organized.get("relation")
                result.append(row)
        return result

    @staticmethod
    def _rows_by_parent(
        rows: list[dict[str, Any]], *, key: str = "parent_record_id"
    ) -> dict[str | None, list[dict[str, Any]]]:
        grouped: dict[str | None, list[dict[str, Any]]] = {}
        for row in rows:
            grouped.setdefault(row.get(key), []).append(row)
        return grouped

    @staticmethod
    def _linked_value(row: dict[str, Any] | None) -> Any:
        if row is None:
            return None
        for field in (
            "value_raw",
            "value_text",
            "value_numeric",
            "value_boolean",
            "value_date",
            "value_code",
        ):
            if row.get(field) is not None:
                return row[field]
        return None

    @staticmethod
    def _headers(rows: list[dict[str, Any]]) -> list[str]:
        keys: list[str] = []
        seen: set[str] = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    keys.append(key)
        return keys

    @staticmethod
    def _cell_value(value: Any) -> Any:
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return json.dumps(value, ensure_ascii=False, sort_keys=True)

    @staticmethod
    def _style_sheet(sheet) -> None:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.alignment = Alignment(vertical="center")
        for column_index, column in enumerate(sheet.columns, 1):
            max_length = max((len(str(cell.value or "")) for cell in column), default=8)
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max_length + 2, 12), 48
            )
            for cell in column:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    def _write_csv(self, path: Path, rows: list[dict[str, Any]]) -> None:
        headers = self._headers(rows)
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers)
            if headers:
                writer.writeheader()
                for row in rows:
                    writer.writerow({key: self._cell_value(row.get(key)) for key in headers})
