from __future__ import annotations

import hashlib
import json
from types import SimpleNamespace

import numpy as np
import pytest
from openpyxl import load_workbook

from esg_targeted.config import Settings
from esg_targeted.contracts import (
    ExtractionRequest,
    ModelRunResult,
    SemanticAssignment,
    SemanticDecision,
    SemanticFactGroup,
)
from esg_targeted.inspection import ResultInspectionService
from esg_targeted.grounding.guard import GroundingGuard
from esg_targeted.models.response_adapter import DirectFillResponseAdapter
from esg_targeted.models.errors import ModelRunFailure
from esg_targeted.models.nuextract import NuExtractMlxModel
from esg_targeted.storage.artifacts import ArtifactStore
from esg_targeted.storage.job_store import JobStore
from esg_targeted.workflow import TargetedExtractionWorkflow
from esg_targeted.results.validation import ResultContractError

from tests.helpers import build_ir_fixture, standard_dist_root
from tests.test_inventory_and_guard import build_pipeline, without_target_cells


def test_keyed_record_merge_is_idempotent_but_rejects_payload_collision():
    records = {"reporting_tasks": []}
    record_index = {
        collection: {}
        for collection in (
            "reporting_tasks",
            "quantitative_observations",
            "qualitative_assertions",
            "attribute_values",
            "dimension_values",
            "evidence_references",
        )
    }
    first = {"reporting_tasks": [{"task_id": "task-1", "task_status": "completed"}]}
    TargetedExtractionWorkflow._merge_records(
        records, record_index, first, "metric-1"
    )
    TargetedExtractionWorkflow._merge_records(
        records, record_index, first, "metric-1"
    )
    assert len(records["reporting_tasks"]) == 1

    with pytest.raises(ResultContractError, match="aggregate collision"):
        TargetedExtractionWorkflow._merge_records(
            records,
            record_index,
            {
                "reporting_tasks": [
                    {"task_id": "task-1", "task_status": "partial"}
                ]
            },
            "metric-2",
        )


def test_checkpoint_recovery_does_not_replace_an_explicit_ambiguous_rerun():
    ambiguous = {
        "schema_version": "targeted-checkpoint-v2",
        "reusable": False,
        "outcome": {"guard_accepted": False, "status": "ambiguous"},
    }
    assert not TargetedExtractionWorkflow._checkpoint_requires_rematerialization(
        ambiguous,
        contract_recovery_job=False,
        package_source_digest="package-digest",
        materializer_version="materializer-v2",
    )
    assert TargetedExtractionWorkflow._checkpoint_requires_rematerialization(
        ambiguous,
        contract_recovery_job=True,
        package_source_digest="package-digest",
        materializer_version="materializer-v2",
    )


class FakeEmbeddingProvider:
    dimensions = 8

    @staticmethod
    def _vector(text: str) -> np.ndarray:
        digest = hashlib.sha256(text.encode("utf-8")).digest()
        vector = np.frombuffer(digest[:8], dtype=np.uint8).astype(np.float32)
        norm = np.linalg.norm(vector)
        return vector / norm if norm else vector

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        return np.stack([self._vector(item) for item in texts])

    def embed_query(self, text: str) -> np.ndarray:
        return self._vector(text)

    def close(self) -> None:
        pass


class NotFoundModel:
    def decide(
        self, packet, *, feedback=None, use_visual=False, cancel_check=None
    ) -> ModelRunResult:
        decision = SemanticDecision(
            task_id=packet.task_id,
            status="not_found",
            fact_groups=[],
            selected_evidence_span_ids=[],
            uncertainty_code="none",
        )
        return ModelRunResult(
            decision=decision,
            raw_output=decision.model_dump_json(),
            cleaned_output=decision.model_dump_json(),
            visual_used=use_visual,
            image_count=0,
        )

    def close(self) -> None:
        pass


class UngroundedModel:
    def decide(
        self, packet, *, feedback=None, use_visual=False, cancel_check=None
    ) -> ModelRunResult:
        span = packet.spans[0]
        decision = SemanticDecision(
            task_id=packet.task_id,
            status="found",
            fact_groups=[
                SemanticFactGroup(
                    group_ref_id=span.context_group_id,
                    assignments=[
                        SemanticAssignment(
                            element_id="esrs.fake.element",
                            source_ref_id=span.span_id,
                            evidence_span_ids=[span.span_id],
                            value_raw=span.text,
                            source_mode="span",
                            confidence=0.8,
                        )
                    ],
                )
            ],
            selected_evidence_span_ids=[span.span_id],
            uncertainty_code="none",
        )
        return ModelRunResult(
            decision=decision,
            raw_output=decision.model_dump_json(),
            cleaned_output=decision.model_dump_json(),
            visual_used=use_visual,
            image_count=0,
        )

    def close(self) -> None:
        pass


class FoundModel:
    def decide(
        self, packet, *, feedback=None, use_visual=False, cancel_check=None
    ) -> ModelRunResult:
        elements = {item["element_code"]: item["element_id"] for item in packet.elements}
        groups = []
        selected = []
        for group_id in packet.allowed_group_ids:
            spans = [item for item in packet.spans if item.context_group_id == group_id]
            pollutant = next(
                (item for item in spans if item.text in {"氮氧化物", "二氧化硫"}),
                None,
            )
            amount = next(
                (
                    item for item in packet.candidates
                    if item.context_group_id == group_id and item.candidate_type == "quantity"
                ),
                None,
            )
            year = next(
                (
                    item for item in packet.candidates
                    if item.context_group_id == group_id and item.candidate_type == "year"
                ),
                None,
            )
            boundary = next(
                (item for item in spans if item.text == "中国境内运营"),
                None,
            )
            if not all((pollutant, amount, year, boundary)):
                continue
            amount_span = next(item for item in spans if item.span_id == amount.span_id)
            assignments = [
                SemanticAssignment(
                    element_id=elements["emission_amount"],
                    source_ref_id=amount.candidate_id,
                    evidence_span_ids=[amount.span_id],
                    value_raw=amount.raw_value,
                    source_mode="candidate",
                    confidence=0.99,
                ),
                SemanticAssignment(
                    element_id=elements["mass_unit"],
                    source_ref_id=amount_span.span_id,
                    evidence_span_ids=[amount_span.span_id],
                    value_raw="吨",
                    source_mode="span",
                    confidence=0.99,
                ),
                SemanticAssignment(
                    element_id=elements["pollutant"],
                    source_ref_id=pollutant.span_id,
                    evidence_span_ids=[pollutant.span_id],
                    value_raw=pollutant.text,
                    source_mode="span",
                    confidence=0.99,
                ),
                SemanticAssignment(
                    element_id=elements["reporting_period"],
                    source_ref_id=year.candidate_id,
                    evidence_span_ids=[year.span_id],
                    value_raw=year.raw_value,
                    source_mode="candidate",
                    confidence=0.99,
                ),
                SemanticAssignment(
                    element_id=elements["reporting_boundary"],
                    source_ref_id=boundary.span_id,
                    evidence_span_ids=[boundary.span_id],
                    value_raw=boundary.text,
                    source_mode="span",
                    confidence=0.99,
                ),
            ]
            groups.append(SemanticFactGroup(group_ref_id=group_id, assignments=assignments))
            selected.extend(
                span_id for item in assignments for span_id in item.evidence_span_ids
            )
        decision = SemanticDecision(
            task_id=packet.task_id,
            status="found" if groups else "not_found",
            fact_groups=groups,
            selected_evidence_span_ids=sorted(set(selected)),
            uncertainty_code="none",
        )
        return ModelRunResult(
            decision=decision,
            raw_output=decision.model_dump_json(),
            cleaned_output=decision.model_dump_json(),
            visual_used=use_visual,
            image_count=0,
        )

    def close(self) -> None:
        pass


class MultiPollutantFoundModel(FoundModel):
    pass


class ContinuationModel:
    provider_name = "continuation_test"
    model_id = "continuation-test-model"

    def __init__(self, first_finish_reason: str = "length") -> None:
        self.calls = 0
        self.feedback: list[str | None] = []
        self.first_finish_reason = first_finish_reason

    def decide(
        self, packet, *, feedback=None, use_visual=False, cancel_check=None
    ) -> ModelRunResult:
        self.calls += 1
        self.feedback.append(feedback)
        context = json.loads(packet.model_context)
        targets = context["region"]["target_value_cells"]
        target = targets[self.calls - 1]
        fields = {item["code"]: None for item in context["elements"]}
        fields["emission_amount"] = target["visible_value"]
        decision, _ = DirectFillResponseAdapter().parse(
            json.dumps(
                {
                    "task_id": packet.task_id,
                    "status": "partial" if self.calls == 1 else "found",
                    "rows": [
                        {
                            "target_cell": target["id"],
                            "group": target["group"],
                            "fields": fields,
                        }
                    ],
                    "uncertainty_code": (
                        "insufficient_evidence" if self.calls == 1 else "none"
                    ),
                },
                ensure_ascii=False,
            ),
            packet=packet,
        )
        return ModelRunResult(
            decision=decision,
            raw_output="{}",
            cleaned_output=decision.model_dump_json(),
            visual_used=use_visual,
            image_count=int(use_visual),
            finish_reason=self.first_finish_reason if self.calls == 1 else "stop",
            output_token_budget=4096,
            provider=self.provider_name,
            model_id=self.model_id,
        )

    def close(self) -> None:
        pass


def _run(tmp_path, model, *, pollutant_rows=None):
    ir_root = tmp_path / "ir-output"
    build_ir_fixture(ir_root, pollutant_rows=pollutant_rows)
    output_root = tmp_path / "outputs"
    settings = Settings(
        project_root=tmp_path,
        ir_roots=(ir_root,),
        standard_dist_root=standard_dist_root(),
        output_root=output_root,
        index_cache_root=tmp_path / "indexes",
        state_db=tmp_path / "jobs.sqlite3",
        nuextract_model_path=tmp_path / "nuextract",
        embedding_model_path=tmp_path / "embedding",
        backend_host="127.0.0.1",
        backend_port=18180,
        frontend_port=18181,
        require_ready_ir=True,
        retrieval_top_k=12,
        packet_max_chars=16000,
        model_max_tokens=512,
        model_retries=1,
    )
    settings.ensure_runtime_dirs()
    jobs = JobStore(settings.state_db)
    artifacts = ArtifactStore(settings.output_root)
    request = ExtractionRequest(
        ir_run_id="ir-fixture-e2-4",
        package_id="esrs.2023-set1.e2-4",
        package_version="1.0.0",
        metric_ids=["esrs.2023-set1.e2-4.dp02"],
    )
    jobs.create("job-fixture", request, str(artifacts.job_dir("job-fixture")))
    workflow = TargetedExtractionWorkflow(
        settings,
        jobs,
        artifacts,
        embedding_provider_factory=FakeEmbeddingProvider,
        semantic_model_factory=lambda: model,
    )
    workflow.run("job-fixture")
    return jobs, artifacts


def test_workflow_writes_all_audit_layers(tmp_path) -> None:
    jobs, artifacts = _run(tmp_path, NotFoundModel())
    record = jobs.get("job-fixture")
    assert record.status.value == "completed"
    paths = {item["path"] for item in artifacts.list("job-fixture")}
    assert {
        "inventory/spans.jsonl",
        "inventory/candidates.jsonl",
        "retrieval/esrs.2023-set1.e2-4.dp02.json",
        "packets/esrs.2023-set1.e2-4.dp02.json",
        "decisions/esrs.2023-set1.e2-4.dp02.json",
        "guards/esrs.2023-set1.e2-4.dp02.json",
        "results/reporting-tasks.jsonl",
        "exports/machine-fill.xlsx",
        "manifest.json",
    } <= paths
    stages = {item["stage"] for item in jobs.events("job-fixture")}
    assert {"semantic_fill", "grounding_guard", "finished"} <= stages


def test_rejected_model_assignment_never_reaches_core_records(tmp_path) -> None:
    jobs, artifacts = _run(tmp_path, UngroundedModel())
    record = jobs.get("job-fixture")
    assert record.status.value == "partial"
    summary = artifacts.read_json("job-fixture", "manifest.json")
    assert summary["guard_accepted_count"] == 0
    assert summary["outcomes"][0]["status"] == "ambiguous"
    result = artifacts.read_json("job-fixture", "exports/extraction-result.json")
    assert result["records"]["quantitative_observations"] == []
    assert result["records"]["qualitative_assertions"] == []


def test_result_bundle_contract_and_read_only_inspection(tmp_path) -> None:
    jobs, artifacts = _run(tmp_path, FoundModel())
    record = jobs.get("job-fixture")
    assert record.status.value == "completed"

    manifest = artifacts.read_json("job-fixture", "manifest.json")
    assert manifest["bundle_layout_version"] == "targeted-result-bundle-v1"
    assert manifest["core_schema"] == {
        "core_schema_id": "extractesg.core",
        "core_schema_version": "1.0.0",
    }
    assert manifest["contract_validation"]["status"] == "passed"
    assert manifest["exports"]["workbook"] == "exports/machine-fill.xlsx"

    package_path = artifacts.resolve("job-fixture", "exports/extraction-result.json")
    before = hashlib.sha256(package_path.read_bytes()).hexdigest()
    inspection = ResultInspectionService(artifacts).build(record)
    after = hashlib.sha256(package_path.read_bytes()).hexdigest()

    assert before == after
    assert inspection.schema_version == "targeted-result-inspection-v1"
    assert inspection.integrity_status == "passed"
    assert inspection.overview.quantitative_fact_count == 1
    assert inspection.overview.evidence_count >= 1
    metric = inspection.metrics[0]
    assert metric.source_datapoint_id == "E2-4_02"
    assert metric.facts[0].value_numeric == "125.6"
    assert metric.facts[0].evidence[0].pdf_page_number == 1
    assert {item.element_code for item in metric.element_columns} == {
        "additional_breakdown",
        "consolidation_scope",
        "emission_amount",
        "mass_unit",
        "pollutant",
        "pollution_medium",
        "reporting_boundary",
        "reporting_period",
        "threshold_basis",
    }
    cells = {item.element_code: item for item in metric.facts[0].element_cells}
    assert cells["emission_amount"].value == "125.6吨"
    assert cells["emission_amount"].provenance_mode == "candidate"
    assert cells["mass_unit"].value == "吨"
    assert cells["mass_unit"].provenance_mode == "span"
    assert cells["pollutant"].value == "氮氧化物"
    assert set(metric.facts[0].provenance_modes) == {"candidate", "span"}
    assert cells["pollution_medium"].value == "air"
    assert cells["pollution_medium"].source == "fixed"
    assert cells["additional_breakdown"].value is None
    assert cells["additional_breakdown"].source == "missing"

    workbook = load_workbook(
        artifacts.resolve("job-fixture", "exports/machine-fill.xlsx"),
        read_only=True,
        data_only=True,
    )
    assert workbook.sheetnames[0] == "机器填写行"
    sheet = workbook["机器填写行"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert [c for c in headers if c in {e.element_code for e in metric.element_columns}] == [e.element_code for e in metric.element_columns]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    machine_row = dict(zip(headers, values))
    assert machine_row["source_datapoint_id"] == "E2-4_02"
    assert machine_row["emission_amount"] == "125.6吨"
    assert machine_row["pollutant"] == "氮氧化物"
    assert machine_row["pollution_medium"] == "air"
    assert machine_row["additional_breakdown"] is None
    assert machine_row["evidence_pdf_pages"] == "[1]"
    workbook.close()


def test_result_inspection_preserves_multiple_facts_for_one_metric(tmp_path) -> None:
    jobs, artifacts = _run(
        tmp_path,
        MultiPollutantFoundModel(),
        pollutant_rows=[
            ("氮氧化物", "125.6吨", "2024年", "中国境内运营"),
            ("二氧化硫", "64.2吨", "2024年", "中国境内运营"),
        ],
    )
    inspection = ResultInspectionService(artifacts).build(jobs.get("job-fixture"))

    assert inspection.overview.quantitative_fact_count == 2
    assert len(inspection.metrics) == 1
    assert len(inspection.metrics[0].facts) == 2
    pollutants = {
        value.value
        for fact in inspection.metrics[0].facts
        for value in fact.dimensions
        if value.label.get("zh") == "污染物"
    }
    assert pollutants == {"氮氧化物", "二氧化硫"}


def _merge_regions(tmp_path, decisions, *, decision_packet=None):
    _, selection, regions, _, _ = build_pipeline(tmp_path)
    packet = decision_packet or regions[0]
    guard_engine = GroundingGuard()
    responses = iter(
        [
            (
                decision,
                guard_engine.validate(packet, decision),
                1,
                [],
            )
            for decision in decisions
        ]
    )
    workflow = object.__new__(TargetedExtractionWorkflow)
    workflow._decide = lambda **_: next(responses)
    return workflow._decide_regions(
        job_id="job-merge",
        metric_id=selection.query.metric_id,
        packet=selection,
        decision_packets=[packet for _ in decisions],
        semantic_model=object(),
        guard_engine=guard_engine,
        request=ExtractionRequest(
            ir_run_id="ir-fixture-e2-4",
            package_id="esrs.2023-set1.e2-4",
            package_version="1.0.0",
        ),
    )


def test_region_merge_found_and_not_found_keeps_found(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)
    found = FoundModel().decide(regions[0]).decision
    not_found = SemanticDecision(
        task_id=regions[0].task_id,
        status="not_found",
        uncertainty_code="none",
    )
    merged, guard, _, _ = _merge_regions(tmp_path, [found, not_found])
    assert merged.status == "found"
    assert merged.fact_groups
    assert guard.accepted


def test_region_merge_deduplicates_identical_rows(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)
    found = FoundModel().decide(regions[0]).decision
    merged, guard, _, _ = _merge_regions(tmp_path, [found, found])
    assert len(merged.fact_groups) == len(found.fact_groups)
    assert guard.accepted


def test_region_merge_partial_and_found_retains_facts_as_partial(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)
    found = FoundModel().decide(regions[0]).decision
    partial = found.model_copy(
        update={
            "status": "partial",
            "uncertainty_code": "insufficient_evidence",
        }
    )
    merged, guard, _, _ = _merge_regions(tmp_path, [partial, found])
    assert merged.status == "partial"
    assert merged.fact_groups
    assert guard.accepted


def test_region_merge_preserves_visual_source_provenance(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)
    packet = without_target_cells(regions[0])
    group_alias = next(iter(packet.alias_map["groups"]))
    fields = {
        item["code"]: None for item in json.loads(packet.model_context)["elements"]
    }
    for field, value in {
        "emission_amount": "777.7",
        "mass_unit": "公斤",
        "pollutant": "二氧化硫",
        "reporting_period": "2023",
        "reporting_boundary": "海外运营",
    }.items():
        fields[field] = value
    decision, _ = DirectFillResponseAdapter().parse(
        json.dumps(
            {
                "task_id": packet.task_id,
                "status": "found",
                "rows": [{"group": group_alias, "fields": fields}],
                "uncertainty_code": "none",
            },
            ensure_ascii=False,
            ),
            packet=packet,
            visual_used=True,
        )

    merged, guard, _, _ = _merge_regions(
        tmp_path,
        [decision],
        decision_packet=packet,
    )

    assert merged.status == "found"
    assert guard.accepted, guard.feedback
    assert all(
        item.source_mode == "visual"
        for group in merged.fact_groups
        for item in group.assignments
    )


def test_region_merge_preserves_physical_target_cell_contract(tmp_path) -> None:
    """A region-accepted cell fact must remain valid after parent-level merge."""

    _, _, regions, _, _ = build_pipeline(tmp_path)
    packet = next(
        item
        for item in regions
        if json.loads(item.model_context)["region"]["target_value_cells"]
    )
    context = json.loads(packet.model_context)
    target = context["region"]["target_value_cells"][0]
    fields = {item["code"]: None for item in context["elements"]}
    fields["emission_amount"] = target["visible_value"]
    decision, _ = DirectFillResponseAdapter().parse(
        json.dumps(
            {
                "task_id": packet.task_id,
                "status": "found",
                "rows": [
                    {
                        "target_cell": target["id"],
                        "group": target["group"],
                        "fields": fields,
                    }
                ],
                "uncertainty_code": "none",
            },
            ensure_ascii=False,
        ),
        packet=packet,
    )
    assert decision.fact_groups[0].group_ref_id == target["cell_id"]
    assert GroundingGuard().validate(packet, decision).accepted

    merged, guard, _, _ = _merge_regions(
        tmp_path,
        [decision],
        decision_packet=packet,
    )

    assert merged.status == "found"
    assert merged.fact_groups[0].group_ref_id == target["cell_id"]
    assert guard.accepted, guard.feedback


def test_failure_routes_distinguish_transport_contract_and_model_output() -> None:
    route = TargetedExtractionWorkflow._failure_retry_strategy

    assert route("service") == "transport_retry_same_evidence"
    assert route("timeout") == "transport_retry_same_evidence"
    assert route("syntax") == "constrained_json_rewrite"
    assert route("schema") == "contract_correction"
    assert route("output_budget") == "compact_output_rewrite"
    assert route("request_config") == "stop_local_contract"


def test_transport_failure_feedback_does_not_ask_model_to_correct_semantics() -> None:
    feedback = TargetedExtractionWorkflow._failure_feedback(
        "service", "upstream unavailable"
    )

    assert "Repeat the same complete evidence task" in feedback
    assert "Correct the previous response" not in feedback


def test_output_budget_continuation_keeps_first_rows_and_fills_remaining_cells(
    tmp_path,
) -> None:
    _, _, regions, _, _ = build_pipeline(
        tmp_path,
        value_columns=[("2023", "1.21"), ("2024", "1.04")],
    )
    packet = next(
        item
        for item in regions
        if len(json.loads(item.model_context)["region"]["target_value_cells"]) == 2
    )

    class EventStore:
        def __init__(self) -> None:
            self.records = []

        def add_event(self, *args, **kwargs) -> None:
            self.records.append((args, kwargs))

        @staticmethod
        def is_cancel_requested(job_id: str) -> bool:
            return False

    model = ContinuationModel()
    workflow = object.__new__(TargetedExtractionWorkflow)
    workflow.settings = SimpleNamespace(model_retries=2)
    workflow.job_store = EventStore()
    decision, guard, attempts, records = workflow._decide(
        job_id="job-continuation",
        metric_id=packet.query.metric_id,
        packet=packet,
        semantic_model=model,
        guard_engine=GroundingGuard(),
        request=ExtractionRequest(
            ir_run_id="ir-fixture-e2-4",
            package_id="esrs.2023-set1.e2-4",
            package_version="1.0.0",
            visual_fallback=False,
        ),
    )

    assert model.calls == 2
    assert attempts == 2
    assert model.feedback[0] is None
    assert "not-yet-covered target cells" in (model.feedback[1] or "")
    assert decision.status == "found"
    assert len(decision.fact_groups) == 2
    assert guard.accepted, guard.feedback
    assert any(item.get("route") == "output_continuation_merge" for item in records)


def test_normal_completion_does_not_force_unrelated_target_coverage(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(
        tmp_path,
        value_columns=[("2023", "1.21"), ("2024", "1.04")],
    )
    packet = next(
        item
        for item in regions
        if len(json.loads(item.model_context)["region"]["target_value_cells"]) == 2
    )

    class EventStore:
        @staticmethod
        def add_event(*args, **kwargs) -> None:
            pass

        @staticmethod
        def is_cancel_requested(job_id: str) -> bool:
            return False

    model = ContinuationModel(first_finish_reason="stop")
    workflow = object.__new__(TargetedExtractionWorkflow)
    workflow.settings = SimpleNamespace(model_retries=2)
    workflow.job_store = EventStore()
    decision, guard, attempts, _ = workflow._decide(
        job_id="job-coverage",
        metric_id=packet.query.metric_id,
        packet=packet,
        semantic_model=model,
        guard_engine=GroundingGuard(),
        request=ExtractionRequest(
            ir_run_id="ir-fixture-e2-4",
            package_id="esrs.2023-set1.e2-4",
            package_version="1.0.0",
            visual_fallback=True,
        ),
    )

    assert model.calls == 1
    assert attempts == 1
    assert len(decision.fact_groups) == 1
    assert guard.accepted, guard.feedback


def test_local_visual_call_receives_full_configured_output_ceiling(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)
    packet = next(item for item in regions if item.page_image_paths)
    model = NuExtractMlxModel(tmp_path / "model", max_tokens=32_768)

    assert model._output_token_budget(packet, visual_used=True) == 32_768
    assert model._output_token_budget(packet, visual_used=False) < 32_768


def test_identical_invalid_output_stops_retries_without_hashing(tmp_path) -> None:
    _, _, regions, _, _ = build_pipeline(tmp_path)

    class RepeatingInvalidModel:
        calls = 0

        def decide(self, *args, **kwargs):
            self.calls += 1
            raise ModelRunFailure("invalid group", category="schema", raw_output='{"group":["G1","G2"]}')

    model = RepeatingInvalidModel()
    workflow = object.__new__(TargetedExtractionWorkflow)
    workflow.settings = SimpleNamespace(model_retries=5)
    workflow.job_store = SimpleNamespace(
        add_event=lambda *args, **kwargs: None,
        is_cancel_requested=lambda job_id: False,
    )
    _, guard, attempts, records = workflow._decide(
        job_id="job-repeated", metric_id=regions[0].query.metric_id, packet=regions[0],
        semantic_model=model, guard_engine=GroundingGuard(),
        request=ExtractionRequest(
            ir_run_id="ir-fixture-e2-4", package_id="esrs.2023-set1.e2-4", package_version="1.0.0",
        ),
    )
    assert model.calls == attempts == 2
    assert not guard.accepted
    assert records[-1]["retry_stopped"] == "identical_contract_output"
